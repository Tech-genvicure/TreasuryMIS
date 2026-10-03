import os
import shutil
import subprocess
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from dotenv import load_dotenv

from .db import get_conn

load_dotenv()

BACKUP_PATH = Path(os.getenv("BACKUP_PATH", "backups"))
RETENTION_DAYS = int(os.getenv("BACKUP_RETENTION_DAYS", "7"))

EXPORT_HEADERS = {
    "Funds_Position": [
        "Name of Company", "Bank Name ", "Account No.", "Balance (INR)",
        "Balance (USD)", "INR Equivalent", "Entity Type"
    ],
    "Payables": [
        "Tracker Serial No.", "Cost For Project", "Status Payment processed [Y/N]",
        "Payment Cycle", "Payer Company", "Branch",
        "Approved IN Vendor Management Software?", "Approval Date",
        "Vendor Name/ Expense Name", "Product/Service", "Whether Pass Through Bill",
        "Against Invoice Ref.", "Invoice Date", "Due Date",
        "Transaction Posting Date", "Amount Payable", "Due Date2",
        "Payment Status Paid/ Not paid", "Bank Account Number", "IFSC CODE",
        "Bank Branch", "Remarks", "Timeline_Bucket"
    ],
    "Borrowinig": ["Sanctioned", "Disbersed", "Outstanding"],
    "Receivables": [
        "hoi", "Cost For Project", "Status Payment processed [Y/N]",
        "Receipt Cycle", "Receiver Company", "Branch",
        "Approved IN Vendor Management Software?", "Approval Date",
        "Vendor Name/ Income Name", "Product/Service", "Whether Pass Through Bill",
        "Against Invoice Ref.", "Invoice Date", "Due Date",
        "Transaction Posting Date", "Amount Receivable", "Due Date2",
        "Payment Status Paid/ Not paid", "Bank Account Number", "IFSC CODE",
        "Bank Branch", "Remarks", "Timline_Bucket"
    ],
}

def _rows(table, columns):
    with get_conn() as conn:
        return conn.execute(
            f"SELECT {', '.join(columns)} FROM {table} ORDER BY id"
        ).fetchall()

def export_excel():
    wb = Workbook()
    wb.remove(wb.active)

    definitions = [
        ("Funds_Position", "funds_position", [
            "company_name", "bank_name", "account_no", "balance_inr",
            "balance_usd", "inr_equivalent", "entity_type"
        ]),
        ("Payables", "payables", [
            "tracker_serial_no", "cost_for_project", "status_payment_processed",
            "payment_cycle", "payer_company", "branch",
            "approved_vendor_management", "approval_date", "vendor_name",
            "product_service", "pass_through_bill", "invoice_ref", "invoice_date",
            "due_date", "transaction_posting_date", "amount_payable", "due_date2",
            "payment_status", "bank_account_number", "ifsc_code", "bank_branch",
            "remarks", "timeline_bucket"
        ]),
        ("Borrowinig", "borrowings", ["sanctioned", "disbursed", "outstanding"]),
        ("Receivables", "receivables", [
            "source_ref", "cost_for_project", "status_payment_processed",
            "receipt_cycle", "receiver_company", "branch",
            "approved_vendor_management", "approval_date", "vendor_name",
            "product_service", "pass_through_bill", "invoice_ref", "invoice_date",
            "due_date", "transaction_posting_date", "amount_receivable", "due_date2",
            "payment_status", "bank_account_number", "ifsc_code", "bank_branch",
            "remarks", "timeline_bucket"
        ]),
    ]

    for sheet_name, table, columns in definitions:
        ws = wb.create_sheet(sheet_name)
        headers = EXPORT_HEADERS[sheet_name]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="123B63")

        for r, row in enumerate(_rows(table, columns), 2):
            for c, col in enumerate(columns, 1):
                ws.cell(r, c, row[col])

        ws.freeze_panes = "A2"
        for column_cells in ws.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 32)
            ws.column_dimensions[column_cells[0].column_letter].width = width

    filename = f"Treasury_Master_{date.today():%d-%b-%Y}.xlsx"
    output = BACKUP_PATH.parent / filename
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output

def backup_database():
    BACKUP_PATH.mkdir(parents=True, exist_ok=True)
    folder = BACKUP_PATH / date.today().isoformat()
    folder.mkdir(parents=True, exist_ok=True)

    dump_file = folder / "treasury_mis.dump"
    database_url = os.getenv("DATABASE_URL")
    pg_dump = os.getenv("PG_DUMP_PATH") or shutil.which("pg_dump")
    if not pg_dump:
        raise RuntimeError("pg_dump was not found. Add PostgreSQL 17 bin to PATH or set PG_DUMP_PATH.")

    subprocess.run(
        [pg_dump, "--format=custom", "--file", str(dump_file), "--dbname", database_url],
        check=True,
    )

    excel_file = export_excel()
    shutil.copy2(excel_file, folder / excel_file.name)

    cutoff = date.today() - timedelta(days=RETENTION_DAYS - 1)
    for child in BACKUP_PATH.iterdir():
        if child.is_dir():
            try:
                folder_date = date.fromisoformat(child.name)
                if folder_date < cutoff:
                    shutil.rmtree(child)
            except ValueError:
                pass

    return folder
