import argparse
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.auth import hash_password
from app.db import get_conn, init_db


def seed_admin(username, password):
    with get_conn() as conn:
        exists = conn.execute(
            "SELECT 1 FROM users WHERE username=%s",
            (username,)
        ).fetchone()

        if exists:
            print(f"Admin '{username}' already exists.")
            return

        conn.execute(
            """
            INSERT INTO users
            (username, password_hash, role)
            VALUES (%s,%s,%s)
            """,
            (
                username,
                hash_password(password),
                "ADMIN",
            ),
        )

        conn.commit()

        print(f"Created admin '{username}'.")


def clean(value):
    """
    Clean values coming from Excel.

    - Blank cells -> None
    - Spaces / whitespace -> None
    - Excel datetime -> date
    - Everything else -> unchanged
    """

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, str):
        value = value.strip()

        if value == "":
            return None

        return value

    return value


def clean_number(value):
    """
    Clean numeric Excel values.

    Blank / whitespace cells become None.
    Numeric strings are converted to numbers.
    """

    value = clean(value)

    if value is None:
        return None

    if isinstance(value, (int, float)):
        return value

    try:
        return float(value)
    except (ValueError, TypeError):
        return None


def import_excel(path, reset=False):

    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(
            f"Excel file not found: {path}"
        )

    print(f"Reading Excel: {path}")

    wb = load_workbook(
        path,
        data_only=True
    )

    # ---------------------------------------------------------
    # RESET EXISTING TREASURY DATA
    # ---------------------------------------------------------

    if reset:

        print("Resetting existing treasury data...")

        with get_conn() as conn:

            conn.execute(
                """
                TRUNCATE TABLE
                    funds_position,
                    payables,
                    receivables,
                    borrowings
                RESTART IDENTITY CASCADE
                """
            )

            conn.commit()

        print("Existing treasury data cleared.")

    # ---------------------------------------------------------
    # IMPORT
    # ---------------------------------------------------------

    with get_conn() as conn:

        # =====================================================
        # FUNDS POSITION
        # =====================================================

        funds_count = 0

        for row in wb["Funds_Position"].iter_rows(
            min_row=2,
            values_only=True
        ):

            if not any(clean(v) is not None for v in row):
                continue

            values = list(row[:7])

            conn.execute(
                """
                INSERT INTO funds_position
                (
                    company_name,
                    bank_name,
                    account_no,
                    balance_inr,
                    balance_usd,
                    inr_equivalent,
                    entity_type
                )
                VALUES (%s,%s,%s,%s,%s,%s,%s)
                """,
                (
                    clean(values[0]),
                    clean(values[1]),
                    clean(values[2]),
                    clean_number(values[3]),
                    clean_number(values[4]),
                    clean_number(values[5]),
                    clean(values[6]),
                ),
            )

            funds_count += 1

        print(f"Funds Position imported: {funds_count}")


        # =====================================================
        # PAYABLES
        # =====================================================

        payables_count = 0

        for row in wb["Payables"].iter_rows(
            min_row=2,
            values_only=True
        ):

            if not any(clean(v) is not None for v in row):
                continue

            values = list(row[:23])

            conn.execute(
                """
                INSERT INTO payables
                (
                    tracker_serial_no,
                    cost_for_project,
                    status_payment_processed,
                    payment_cycle,
                    payer_company,
                    branch,
                    approved_vendor_management,
                    approval_date,
                    vendor_name,
                    product_service,
                    pass_through_bill,
                    invoice_ref,
                    invoice_date,
                    due_date,
                    transaction_posting_date,
                    amount_payable,
                    due_date2,
                    payment_status,
                    bank_account_number,
                    ifsc_code,
                    bank_branch,
                    remarks,
                    timeline_bucket
                )
                VALUES
                (
                    %s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,
                    %s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s
                )
                """,
                (
                    clean(values[0]),
                    clean(values[1]),
                    clean(values[2]),
                    clean(values[3]),
                    clean(values[4]),
                    clean(values[5]),
                    clean(values[6]),
                    clean(values[7]),
                    clean(values[8]),
                    clean(values[9]),
                    clean(values[10]),
                    clean(values[11]),
                    clean(values[12]),
                    clean(values[13]),
                    clean(values[14]),
                    clean_number(values[15]),
                    clean(values[16]),
                    clean(values[17]),
                    clean(values[18]),
                    clean(values[19]),
                    clean(values[20]),
                    clean(values[21]),
                    clean(values[22]),
                ),
            )

            payables_count += 1

        print(f"Payables imported: {payables_count}")


        # =====================================================
        # RECEIVABLES
        # =====================================================

        receivables_count = 0

        for row in wb["Receivables"].iter_rows(
            min_row=2,
            values_only=True
        ):

            if not any(clean(v) is not None for v in row):
                continue

            values = list(row[:23])

            conn.execute(
                """
                INSERT INTO receivables
                (
                    source_ref,
                    cost_for_project,
                    status_payment_processed,
                    receipt_cycle,
                    receiver_company,
                    branch,
                    approved_vendor_management,
                    approval_date,
                    vendor_name,
                    product_service,
                    pass_through_bill,
                    invoice_ref,
                    invoice_date,
                    due_date,
                    transaction_posting_date,
                    amount_receivable,
                    due_date2,
                    payment_status,
                    bank_account_number,
                    ifsc_code,
                    bank_branch,
                    remarks,
                    timeline_bucket
                )
                VALUES
                (
                    %s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,
                    %s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s
                )
                """,
                (
                    clean(values[0]),
                    clean(values[1]),
                    clean(values[2]),
                    clean(values[3]),
                    clean(values[4]),
                    clean(values[5]),
                    clean(values[6]),
                    clean(values[7]),
                    clean(values[8]),
                    clean(values[9]),
                    clean(values[10]),
                    clean(values[11]),
                    clean(values[12]),
                    clean(values[13]),
                    clean(values[14]),
                    clean_number(values[15]),
                    clean(values[16]),
                    clean(values[17]),
                    clean(values[18]),
                    clean(values[19]),
                    clean(values[20]),
                    clean(values[21]),
                    clean(values[22]),
                ),
            )

            receivables_count += 1

        print(f"Receivables imported: {receivables_count}")


        # =====================================================
        # BORROWINGS
        # =====================================================

        borrow = next(
            wb["Borrowinig"].iter_rows(
                min_row=2,
                values_only=True
            ),
            None
        )

        if borrow and any(
            clean(v) is not None
            for v in borrow
        ):

            conn.execute(
                "DELETE FROM borrowings"
            )

            conn.execute(
                """
                INSERT INTO borrowings
                (
                    sanctioned,
                    disbursed,
                    outstanding
                )
                VALUES (%s,%s,%s)
                """,
                (
                    clean_number(borrow[0]) or 0,
                    clean_number(borrow[1]) or 0,
                    clean_number(borrow[2]) or 0,
                ),
            )

            print("Borrowing position imported.")

        else:
            print("No borrowing data found.")

        conn.commit()

    print()
    print("=" * 55)
    print("Excel import completed successfully.")
    print("=" * 55)
    print(f"File: {path.resolve()}")


if __name__ == "__main__":

    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--username",
        default="admin"
    )

    parser.add_argument(
        "--password",
        required=True
    )

    parser.add_argument(
        "--excel"
    )

    parser.add_argument(
        "--reset",
        action="store_true"
    )

    args = parser.parse_args()

    # Create tables
    init_db()

    # Create admin if required
    seed_admin(
        args.username,
        args.password
    )

    # Import Excel
    if args.excel:

        import_excel(
            args.excel,
            args.reset
        )